"""to_blocks(): a DoclingDocument becomes contract Blocks.

Part 1 builds small DoclingDocuments by hand, so each rule is tested on exactly the case it is about.
Part 2 runs real Docling on a Word file and a PowerPoint file generated here (no layout model: a few seconds).
"""
from pathlib import Path

import pytest
from docling_core.types.doc import (BoundingBox, ContentLayer, DocItemLabel, DoclingDocument, GroupLabel,
                                    ProvenanceItem, TableData)

from app.ingestion.contract import Block, Location, ParsedDocument, QualityReport, SourceType
from app.ingestion.docling_blocks import to_blocks

TITLE = "DSA notes"


def _prov(page: int) -> ProvenanceItem:
    return ProvenanceItem(page_no=page, bbox=BoundingBox(l=0, t=0, r=1, b=1), charspan=(0, 0))


def _blocks(doc: DoclingDocument, source_type: SourceType = "pdf") -> list[Block]:
    """to_blocks(), plus the contract check every loader runs on its result."""
    blocks = to_blocks(doc, source_type, TITLE)
    ParsedDocument(resource_id="r", course_id="c", file_name="notes", file_sha256="0" * 64,
                   source_type=source_type, title=TITLE, parser="test", parser_version="1", blocks=blocks,
                   quality=QualityReport(status="ok", units_total=1, chars_total=1))
    return blocks


# ================================================================ part 1: hand-built documents

def test_headings_build_the_path_and_a_new_heading_closes_the_old_section():
    doc = DoclingDocument(name="t")
    doc.add_heading("Linked Lists", level=1, prov=_prov(1))
    doc.add_text(DocItemLabel.TEXT, "A list stores values in nodes.", prov=_prov(1))
    doc.add_heading("Singly linked", level=2, prov=_prov(2))
    doc.add_text(DocItemLabel.TEXT, "Each node points to the next.", prov=_prov(2))
    doc.add_heading("Stacks", level=1, prov=_prov(3))
    doc.add_text(DocItemLabel.TEXT, "Last in, first out.", prov=_prov(3))
    assert [b.heading_path for b in _blocks(doc)] == [
        [TITLE, "Linked Lists"], [TITLE, "Linked Lists"],
        [TITLE, "Linked Lists", "Singly linked"], [TITLE, "Linked Lists", "Singly linked"],
        [TITLE, "Stacks"], [TITLE, "Stacks"],
    ]


def test_a_title_goes_above_every_heading_but_the_document_title_is_not_repeated():
    doc = DoclingDocument(name="t")
    doc.add_title("dsa  NOTES", prov=_prov(1))                        # the same words as TITLE
    doc.add_title("Chapter 3: Lists", prov=_prov(2))
    doc.add_heading("3.1 Nodes", level=1, prov=_prov(2))
    doc.add_text(DocItemLabel.TEXT, "A node holds a value.", prov=_prov(2))
    blocks = _blocks(doc)
    assert [b.text for b in blocks] == ["Chapter 3: Lists", "3.1 Nodes", "A node holds a value."]
    assert blocks[-1].heading_path == [TITLE, "Chapter 3: Lists", "3.1 Nodes"]


def test_a_pdf_block_gets_its_page_and_a_page_range_when_it_runs_on():
    doc = DoclingDocument(name="t")
    doc.add_text(DocItemLabel.TEXT, "Fits on page four.", prov=_prov(4))
    long_one = doc.add_text(DocItemLabel.TEXT, "Starts on four and ends on five.", prov=_prov(4))
    long_one.prov.append(_prov(5))
    assert [(b.loc.page, b.loc.page_end) for b in _blocks(doc)] == [(4, None), (4, 5)]


def test_headers_footers_and_contents_pages_are_left_out():
    doc = DoclingDocument(name="t")
    doc.add_text(DocItemLabel.PAGE_HEADER, "CS-201 Data Structures", prov=_prov(1),
                 content_layer=ContentLayer.FURNITURE)
    doc.add_text(DocItemLabel.PAGE_FOOTER, "Chapter 3 | 41", prov=_prov(1))    # labelled a footer inside the body
    doc.add_table(data=TableData(num_rows=0, num_cols=0), label=DocItemLabel.DOCUMENT_INDEX, prov=_prov(1))
    doc.add_text(DocItemLabel.TEXT, "The only real content.", prov=_prov(1))
    assert [b.text for b in _blocks(doc)] == ["The only real content."]


def test_words_inside_a_figure_become_figure_text_and_its_caption_stays_a_caption():
    doc = DoclingDocument(name="t")
    doc.add_heading("Linked Lists", level=1, prov=_prov(2))
    picture = doc.add_picture(prov=_prov(2))
    doc.add_text(DocItemLabel.TEXT, "head", parent=picture, prov=_prov(2))
    doc.add_heading("next", level=1, parent=picture, prov=_prov(2))     # big text in a diagram is not a heading
    caption = doc.add_text(DocItemLabel.CAPTION, "Figure 3.1: A list of three nodes", parent=picture,
                           prov=_prov(2))
    picture.captions.append(caption.get_ref())
    blocks = _blocks(doc)
    assert [(b.kind, b.text) for b in blocks] == [
        ("heading", "Linked Lists"), ("figure_text", "head"), ("figure_text", "next"),
        ("caption", "Figure 3.1: A list of three nodes"),
    ]
    assert all(b.heading_path == [TITLE, "Linked Lists"] for b in blocks)


def test_code_keeps_its_spacing_and_list_items_keep_their_markers():
    doc = DoclingDocument(name="t")
    code = "for (int i = 0; i < n; i++)\n    sum += a[i];"
    doc.add_code(code, prov=_prov(1))
    numbered = doc.add_list_group()
    doc.add_list_item("Create the node", enumerated=True, marker="1.", parent=numbered, prov=_prov(1))
    bullets = doc.add_list_group()
    doc.add_list_item("Insert at head", parent=bullets, prov=_prov(1))
    assert [(b.kind, b.text, b.lang) for b in _blocks(doc)] == [
        ("code", code, "code"), ("list_item", "1. Create the node", "en"), ("list_item", "- Insert at head", "en"),
    ]


def test_speaker_notes_are_kept_and_reviewers_comments_are_not():
    doc = DoclingDocument(name="t")
    slide = doc.add_group(label=GroupLabel.SLIDE, name="slide-1")
    doc.add_title("Stacks", parent=slide, prov=_prov(1))
    doc.add_text(DocItemLabel.TEXT, "Mention the call stack.", parent=slide, prov=_prov(1),
                 content_layer=ContentLayer.NOTES)
    comments = doc.add_group(label=GroupLabel.COMMENT_SECTION, name="comment-1", content_layer=ContentLayer.NOTES)
    doc.add_comment(text="[author: Areeba]: fix this slide", parent=comments)
    assert [(b.kind, b.text, b.loc.slide) for b in _blocks(doc, "pptx")] == [
        ("heading", "Stacks", 1), ("slide_notes", "Mention the call stack.", 1),
    ]


def test_empty_texts_and_stray_page_numbers_are_dropped_without_gaps_in_seq():
    doc = DoclingDocument(name="t")
    doc.add_text(DocItemLabel.TEXT, "First.", prov=_prov(1))
    doc.add_text(DocItemLabel.TEXT, "   ", prov=_prov(1))
    doc.add_text(DocItemLabel.TEXT, "41", prov=_prov(1))                 # a page number labelled as text
    doc.add_code("\n\n", prov=_prov(1))
    doc.add_text(DocItemLabel.TEXT, "Second.", prov=_prov(2))
    assert [(b.seq, b.text) for b in _blocks(doc)] == [(0, "First."), (1, "Second.")]


def test_formats_docling_does_not_read_are_refused():
    with pytest.raises(ValueError):
        to_blocks(DoclingDocument(name="t"), "audio", TITLE)


# ================================================================ part 2: real Docling on generated files

def _write_word_file(path: Path) -> Path:
    from docx import Document

    d = Document()
    d.add_heading("Data Structures Notes", level=0)
    d.add_heading("Linked Lists", level=1)
    d.add_paragraph("A linked list stores each value in its own node.")
    d.add_heading("Singly linked lists", level=2)
    d.add_paragraph("Each node keeps a pointer to the next node.")
    d.add_paragraph("Insert at head", style="List Bullet")
    d.add_paragraph("Create the node", style="List Number")
    table = d.add_table(rows=2, cols=2)
    for col, (top, bottom) in enumerate([("Operation", "Insert at head"), ("Cost", "O(1)")]):
        table.cell(0, col).text, table.cell(1, col).text = top, bottom
    d.add_heading("Stacks", level=1)
    code = d.add_paragraph("void push(int x) { top = new Node(x, top); }")
    for run in code.runs:
        run.font.name = "Courier New"                                    # Docling reads monospace text as code
    d.sections[0].header.paragraphs[0].text = "CS-201 Data Structures"
    d.sections[0].footer.paragraphs[0].text = "Page footer text"
    d.save(path)
    return path


def _write_slides(path: Path) -> Path:
    from pptx import Presentation

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])                  # title and content
    first.shapes.title.text = "Stacks"
    first.placeholders[1].text_frame.text = "Last in, first out"
    first.notes_slide.notes_text_frame.text = "Mention the call stack as a real example."
    second = deck.slides.add_slide(deck.slide_layouts[5])                 # title only
    second.shapes.title.text = "Costs"
    deck.save(path)
    return path


@pytest.fixture(scope="module")
def converter():
    from docling.document_converter import DocumentConverter            # imported here: unit tests stay fast

    return DocumentConverter()


@pytest.fixture(scope="module")
def word_blocks(converter, tmp_path_factory) -> list[Block]:
    path = _write_word_file(tmp_path_factory.mktemp("word") / "notes.docx")
    return _blocks(converter.convert(path).document, "docx")


@pytest.fixture(scope="module")
def slide_blocks(converter, tmp_path_factory) -> list[Block]:
    path = _write_slides(tmp_path_factory.mktemp("slides") / "notes.pptx")
    return _blocks(converter.convert(path).document, "pptx")


def test_a_word_file_becomes_blocks_in_reading_order(word_blocks):
    assert [(b.kind, b.text) for b in word_blocks] == [
        ("heading", "Data Structures Notes"),
        ("heading", "Linked Lists"),
        ("paragraph", "A linked list stores each value in its own node."),
        ("heading", "Singly linked lists"),
        ("paragraph", "Each node keeps a pointer to the next node."),
        ("list_item", "- Insert at head"),
        ("list_item", "1. Create the node"),
        ("table", word_blocks[7].text),
        ("heading", "Stacks"),
        ("code", "void push(int x) { top = new Node(x, top); }"),
    ]
    assert "| Insert at head | O(1)" in word_blocks[7].text                  # the table, as Markdown
    assert word_blocks[-1].heading_path == [TITLE, "Data Structures Notes", "Stacks"]


def test_a_word_file_has_no_page_numbers_and_no_header_or_footer(word_blocks):
    assert all(b.loc == Location() for b in word_blocks)
    assert not any("CS-201" in b.text or "footer" in b.text for b in word_blocks)


def test_slides_cite_their_number_and_keep_speaker_notes(slide_blocks):
    assert [(b.kind, b.text, b.loc.slide) for b in slide_blocks] == [
        ("heading", "Stacks", 1),
        ("list_item", "- Last in, first out", 1),
        ("slide_notes", "Mention the call stack as a real example.", 1),
        ("heading", "Costs", 2),
    ]
    assert slide_blocks[-1].heading_path == [TITLE, "Costs"]