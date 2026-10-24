import shutil
from collections import Counter

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
import os
from PIL import Image, ImageDraw, ImageFont
from app.ingestion import messages
from app.ingestion.contract import Block, Location
from app.ingestion.loaders.docx import (OCR_HEADER, OcrLine, _add_picture_text, _drop_word_extras,
                                        _mark_captions, _mark_code, _picture_text, _promote_titles)
from app.ingestion.pipeline import process

@pytest.fixture
def make_docx(tmp_path):
    def make(name="week1.docx", build=None):
        doc = Document()
        if build:
            build(doc)
        path = tmp_path / name
        doc.save(path)
        return path
    return make


def _lecture(doc):
    doc.add_heading("Intro", 1)
    doc.add_paragraph("Hello world paragraph.")
    doc.add_paragraph("Figure 1: a diagram", style="Caption")
    doc.add_heading("Details", 2)
    doc.add_paragraph("More text here.")


def _block(text, kind="paragraph"):
    return Block(seq=0, kind=kind, text=text, heading_path=["t"], loc=Location(), lang="en")


def test_a_word_file_goes_all_the_way_to_chunks_json(make_docx, data_dir, word_tokens):
    result = process(make_docx(build=_lecture), "w1", "c1")
    assert (result.status, result.warnings) == ("ok", [])
    assert (data_dir / "w1" / "parsed.json").exists() and (data_dir / "w1" / "chunks.json").exists()
    assert result.chunked.chunker == "docx-sections/1"
    assert all(b.heading_path[0] == result.parsed.title for b in result.parsed.blocks)
    assert all(b.loc == Location() for b in result.parsed.blocks)


def test_a_caption_paragraph_becomes_a_caption_block(make_docx, data_dir, word_tokens):
    result = process(make_docx(build=_lecture), "w2", "c1")
    captions = [b for b in result.parsed.blocks if b.kind == "caption"]
    assert [b.text for b in captions] == ["Figure 1: a diagram"]


def test_header_and_footer_stay_out(make_docx, data_dir, word_tokens):
    def build(doc):
        doc.sections[0].header.paragraphs[0].text = "RUNNING HEADER"
        doc.sections[0].footer.paragraphs[0].text = "RUNNING FOOTER"
        _lecture(doc)

    result = process(make_docx(build=build), "w3", "c1")
    text = " ".join(b.text for b in result.parsed.blocks)
    assert "RUNNING HEADER" not in text and "RUNNING FOOTER" not in text


def test_a_table_becomes_one_markdown_block(make_docx, data_dir, word_tokens):
    def build(doc):
        doc.add_heading("Data", 1)
        table = doc.add_table(rows=2, cols=2)
        for cell, value in zip(table._cells, ["Key", "Value", "a", "1"]):
            cell.text = value
        doc.add_paragraph("After the table.")

    result = process(make_docx(build=build), "w4", "c1")
    tables = [b for b in result.parsed.blocks if b.kind == "table"]
    assert len(tables) == 1 and "|" in tables[0].text and "Key" in tables[0].text


def test_each_heading_section_gets_its_own_chunks(make_docx, data_dir, word_tokens):
    result = process(make_docx(build=_lecture), "w5", "c1")
    paths = [tuple(c.heading_path) for c in result.chunked.chunks]
    assert len(paths) == 2 and paths[0] != paths[1]
    assert paths[1][-1] == "Details"


def test_a_file_without_text_fails_politely(make_docx, data_dir, word_tokens):
    result = process(make_docx(name="empty.docx"), "w6", "c1")
    assert (result.status, result.warnings, result.chunked) == ("failed", [messages.NO_TEXT_FILE], None)
    assert not (data_dir / "w6" / "chunks.json").exists()


def test_a_broken_file_fails_without_chunks(tmp_path, data_dir, word_tokens):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04 this is not a real zip")
    result = process(path, "w7", "c1")
    assert result.status == "failed" and result.chunked is None
    assert not (data_dir / "w7" / "chunks.json").exists()


def test_a_pdf_renamed_to_docx_is_refused(digital_pdf, tmp_path, data_dir):
    fake = tmp_path / "fake.docx"
    shutil.copy(digital_pdf, fake)
    result = process(fake, "w8", "c1")
    assert (result.status, result.warnings) == ("failed", [messages.TYPE_MISMATCH])


def test_processing_the_same_word_file_twice_gives_the_same_chunks(make_docx, data_dir, word_tokens):
    path = make_docx(build=_lecture)
    process(path, "w9", "c1")
    first = (data_dir / "w9" / "chunks.json").read_bytes()
    process(path, "w9", "c1")
    assert (data_dir / "w9" / "chunks.json").read_bytes() == first


def test_only_paragraphs_with_a_caption_text_are_marked():
    blocks = [_block("Figure 1:  a diagram"), _block("Plain text"), _block("Figure 1: a diagram", "list_item")]
    marked = _mark_captions(blocks, Counter({"Figure 1: a diagram": 1}))
    assert [b.kind for b in marked] == ["caption", "paragraph", "list_item"]


def test_each_caption_text_is_used_once():
    blocks = [_block("Figure 1: a diagram"), _block("Figure 1: a diagram")]
    marked = _mark_captions(blocks, Counter({"Figure 1: a diagram": 1}))
    assert [b.kind for b in marked] == ["caption", "paragraph"]

def _add_toc(doc, entries):
    """A table of contents the way Word stores it: a content control holding the TOC lines."""
    lines = "".join(
        f'<w:p><w:r><w:t>{text}</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>{page}</w:t></w:r></w:p>'
        for text, page in entries
    )
    doc.element.body.insert(0, parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
        f'</w:docPartObj></w:sdtPr><w:sdtContent>'
        f'<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Contents</w:t></w:r></w:p>{lines}'
        f'</w:sdtContent></w:sdt>'
    ))


def test_the_table_of_contents_stays_out(make_docx, data_dir, word_tokens):
    def build(doc):
        _add_toc(doc, [("Intro", 1), ("Details", 2)])
        _lecture(doc)

    result = process(make_docx(build=build), "w10", "c1")
    texts = [b.text for b in result.parsed.blocks]
    assert "Contents" not in texts and "Intro 1" not in texts and "Details 2" not in texts
    assert [b.text for b in result.parsed.blocks if b.kind == "heading"] == ["Intro", "Details"]


def _add_cover_page(doc, title, author):
    """A cover page the way Word stores it: a content control marked "Cover Pages"."""
    doc.element.body.insert(0, parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Cover Pages"/>'
        f'</w:docPartObj></w:sdtPr><w:sdtContent>'
        f'<w:p><w:pPr><w:pStyle w:val="Title"/></w:pPr><w:r><w:t>{title}</w:t></w:r></w:p>'
        f'<w:p><w:r><w:t>{author}</w:t></w:r></w:p>'
        f'</w:sdtContent></w:sdt>'
    ))


def test_the_cover_page_stays_out_of_blocks_and_heading_paths(make_docx, data_dir, word_tokens):
    def build(doc):
        _add_cover_page(doc, "Cover Title", "Cover Author")
        _lecture(doc)

    result = process(make_docx(build=build), "w11", "c1")
    assert not any("Cover" in b.text for b in result.parsed.blocks)
    assert not any("Cover Title" in b.heading_path for b in result.parsed.blocks)


def test_hidden_text_stays_out(make_docx, data_dir, word_tokens):
    def build(doc):
        doc.add_heading("Intro", 1)
        paragraph = doc.add_paragraph("Shown before ")
        paragraph.add_run("SECRET ").font.hidden = True
        paragraph.add_run("shown after.")
        doc.add_paragraph("Whole paragraph hidden.").runs[0].font.hidden = True
        doc.add_paragraph("Last visible paragraph.")

    result = process(make_docx(build=build), "w12", "c1")
    texts = [b.text for b in result.parsed.blocks]
    assert "Shown before shown after." in texts and "Last visible paragraph." in texts
    assert not any("SECRET" in t or "Whole paragraph hidden" in t for t in texts)


def test_a_real_heading_with_the_same_words_as_a_toc_line_is_kept():
    blocks = [_block("Intro"), _block("Intro", "heading"), _block("Body")]
    kept = _drop_word_extras(blocks, Counter({"Intro": 1}))
    assert [(b.kind, b.text) for b in kept] == [("heading", "Intro"), ("paragraph", "Body")]


def test_a_dropped_heading_is_taken_out_of_the_paths_below_it():
    def block(seq, kind, text, path):
        return Block(seq=seq, kind=kind, text=text, heading_path=path, loc=Location(), lang="en")

    blocks = [
        block(0, "heading", "Cover", ["t", "Cover"]),
        block(1, "heading", "Intro", ["t", "Cover", "Intro"]),
        block(2, "paragraph", "Body", ["t", "Cover", "Intro"]),
    ]
    kept = _drop_word_extras(blocks, Counter({"Cover": 1}))
    assert [(b.text, b.heading_path) for b in kept] == [("Intro", ["t", "Intro"]), ("Body", ["t", "Intro"])]
def _commands(doc):
    """Shell commands the way the Bandit notes have them: one Consolas paragraph per line."""
    def code_line(text):
        doc.add_paragraph().add_run(text).font.name = "Consolas"

    doc.add_heading("Level 12", 2)
    doc.add_paragraph("Commands Used:")
    code_line("mktemp -d          # make a work folder")
    code_line("cd /tmp/work")
    code_line("")                                           # a blank line inside the commands
    code_line("ls -la")
    doc.add_paragraph("Explanation follows here.")


def test_monospace_commands_become_one_code_block_with_word_spacing(make_docx, data_dir, word_tokens):
    result = process(make_docx(build=_commands), "w13", "c1")
    code = [b for b in result.parsed.blocks if b.kind == "code"]
    assert [b.text for b in code] == ["mktemp -d          # make a work folder\ncd /tmp/work\n\nls -la"]
    assert code[0].lang == "code"
    texts = [b.text for b in result.parsed.blocks]
    assert "Commands Used:" in texts and "Explanation follows here." in texts


def test_commands_in_a_file_without_a_normal_style_become_code(make_docx, data_dir, word_tokens):
    def build(doc):
        doc.styles["Normal"].element.attrib.pop(qn("w:default"))   # like the Bandit notes: no default style
        _commands(doc)

    result = process(make_docx(build=build), "w14", "c1")
    assert [b.kind for b in result.parsed.blocks].count("code") == 1


def test_a_file_typed_all_in_a_monospace_font_is_not_turned_into_code(make_docx, data_dir, word_tokens):
    def build(doc):
        doc.styles["Normal"].font.name = "Courier New"
        doc.add_heading("Notes", 1)
        doc.add_paragraph("Arrays keep items next to each other in memory.")
        doc.add_paragraph("Linked lists keep a pointer to the next item.")

    result = process(make_docx(build=build), "w15", "c1")
    assert "code" not in [b.kind for b in result.parsed.blocks]
def test_paragraph_and_code_blocks_of_one_run_become_one_code_block():
    blocks = [_block("Commands Used:"), _block("ls -la"), _block('echo "<p>" | nc localhost 30000', "code"),
              _block("After.")]
    marked = _mark_code(blocks, [["ls -la", 'echo "<p>" | nc localhost 30000']])
    assert [(b.kind, b.text) for b in marked] == [
        ("paragraph", "Commands Used:"), ("code", 'ls -la\necho "<p>" | nc localhost 30000'), ("paragraph", "After.")]


def test_code_text_comes_from_word_with_indentation_and_lines_the_cleaner_dropped():
    blocks = [_block("int main() {"), _block("return 0;"), _block("After.")]   # "}" was cleaned away
    marked = _mark_code(blocks, [["int main() {", "    return 0;", "}"]])
    assert [(b.kind, b.text) for b in marked] == [("code", "int main() {\n    return 0;\n}"), ("paragraph", "After.")]
def _bold_line(doc, text, size):
    run = doc.add_paragraph().add_run(text)
    run.bold, run.font.size = True, Pt(size)


def test_bold_titles_become_headings_in_a_file_without_heading_styles(make_docx, data_dir, word_tokens):
    def build(doc):                                         # like Results.docx: no heading style at all
        _bold_line(doc, "Outcomes", 18)
        _bold_line(doc, "Pain score 24 hours", 14)
        doc.add_paragraph("Pain was lower on the first day.")
        _bold_line(doc, "Pain score 48 hours", 14)
        doc.add_paragraph("Pain was the same on the second day.")

    result = process(make_docx(build=build), "w16", "c1")
    headings = [(b.text, b.heading_path[1:]) for b in result.parsed.blocks if b.kind == "heading"]
    assert headings == [("Outcomes", ["Outcomes"]),
                        ("Pain score 24 hours", ["Outcomes", "Pain score 24 hours"]),
                        ("Pain score 48 hours", ["Outcomes", "Pain score 48 hours"])]
    chunks = {c.heading_path[-1]: c.text for c in result.chunked.chunks}
    assert chunks["Pain score 48 hours"] == "Pain was the same on the second day."


def test_bold_lines_stay_paragraphs_when_the_file_has_real_headings(make_docx, data_dir, word_tokens):
    def build(doc):                                         # like the Bandit notes: "Goal:" is only a label
        doc.add_heading("Level 0", 1)
        _bold_line(doc, "Goal:", 11)
        doc.add_paragraph("Log in with ssh.")

    result = process(make_docx(build=build), "w17", "c1")
    assert [b.text for b in result.parsed.blocks if b.kind == "heading"] == ["Level 0"]
    assert "Goal:" in [b.text for b in result.parsed.blocks if b.kind == "paragraph"]
def test_bold_titles_nest_by_font_size_and_set_the_path_of_the_text_below():
    blocks = [_block("T"), _block("Outcomes"), _block("Pain score 48 hours"), _block("Pain was the same."),
              _block("Arrhythmia"), _block("No difference.")]
    titles = [("T", 1), ("Outcomes", 1), ("Pain score 48 hours", 2), ("Arrhythmia", 2)]
    marked = _promote_titles(blocks, titles, "t")           # "T" has the file title's words: it adds no level
    assert [(b.kind, b.heading_path) for b in marked] == [
        ("heading", ["t"]),
        ("heading", ["t", "Outcomes"]),
        ("heading", ["t", "Outcomes", "Pain score 48 hours"]),
        ("paragraph", ["t", "Outcomes", "Pain score 48 hours"]),
        ("heading", ["t", "Outcomes", "Arrhythmia"]),
        ("paragraph", ["t", "Outcomes", "Arrhythmia"]),
    ]
needs_tesseract = pytest.mark.skipif(
    shutil.which("tesseract") is None and not os.path.exists(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
    reason="picture OCR test: install Tesseract (winget install UB-Mannheim.TesseractOCR)",
)


def _text_picture(path, lines):
    """A screenshot-like picture: black text on white."""
    image = Image.new("RGB", (900, 60 + 50 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=32)
    for i, line in enumerate(lines):
        draw.text((20, 20 + 50 * i), line, fill="black", font=font)
    image.save(path)
    return path


def _lecture_with_screenshot(picture):
    def build(doc):
        doc.add_heading("Searching", 1)
        doc.add_paragraph("Before the picture.")
        doc.add_picture(str(picture))
        doc.add_paragraph("After the picture.")
    return build


@needs_tesseract
def test_text_in_a_picture_is_read_and_placed_after_the_paragraph_before_it(make_docx, data_dir, word_tokens,
                                                                           tmp_path):
    picture = _text_picture(tmp_path / "shot.png", ["Binary search needs a sorted array",
                                                    "and runs in logarithmic time"])
    result = process(make_docx(build=_lecture_with_screenshot(picture)), "w18", "c1")
    texts = [b.text for b in result.parsed.blocks]
    i = texts.index("Before the picture.")
    picture_block = result.parsed.blocks[i + 1]
    assert picture_block.kind == "figure_text" and picture_block.heading_path[-1] == "Searching"
    assert "Binary search needs a sorted array" in picture_block.text
    assert texts[i + 2] == "After the picture."
    assert (result.status, result.warnings) == ("ok", [messages.PICTURE_OCR])


@needs_tesseract
def test_picture_ocr_can_be_switched_off(make_docx, data_dir, word_tokens, tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_PICTURE_OCR", "0")
    picture = _text_picture(tmp_path / "shot.png", ["Binary search needs a sorted array"])
    result = process(make_docx(build=_lecture_with_screenshot(picture)), "w19", "c1")
    assert "figure_text" not in [b.kind for b in result.parsed.blocks]
    assert (result.status, result.warnings) == ("ok", [])


def test_picture_text_keeps_indentation_bullets_and_drops_unsure_lines():
    lines = [OcrLine("def f(x):", 10, 10.0, 0.9), OcrLine("if x:", 50, 10.0, 0.9),
             OcrLine("return 1", 90, 10.0, 0.9), OcrLine("~~ noise", 10, 10.0, 0.2),
             OcrLine("¢ a bullet point", 10, 10.0, 0.9)]
    assert _picture_text(lines) == "def f(x):\n    if x:\n        return 1\n- a bullet point"


def test_picture_text_goes_after_the_block_before_it_and_icons_and_logos_are_skipped():
    def block(ref, text):
        return Block(seq=0, kind="paragraph", text=text, heading_path=["t", "Sorting"], loc=Location(),
                     extra={"ref": ref})

    before, after = block("#/texts/0", "Before."), block("#/texts/2", "After.")
    screenshot, icon, logo = Image.new("RGB", (400, 100)), Image.new("RGB", (20, 20)), Image.new("RGB", (300, 200))
    pictures = [("#/pictures/0", screenshot, ["#/texts/0", "#/texts/1"]),   # "#/texts/1" was dropped earlier
                ("#/pictures/1", icon, ["#/texts/0", "#/texts/1"]),         # too small: not read at all
                ("#/pictures/2", logo, ["#/texts/0", "#/texts/1"])]         # scraps, no real words: skipped

    def read(image):
        if image.size == (400, 100):
            return [OcrLine("Merge sort splits the array in half", 0, 10.0, 0.9)]
        return [OcrLine("/ Ve “é", 0, 10.0, 0.5), OcrLine("YZ J", 0, 10.0, 0.5), OcrLine("nl ‘ /", 0, 10.0, 0.5)]

    out, count = _add_picture_text([before, after], pictures, read, "t")
    assert count == 1
    assert [(b.kind, b.heading_path) for b in out] == [
        ("paragraph", ["t", "Sorting"]), ("figure_text", ["t", "Sorting"]), ("paragraph", ["t", "Sorting"])]
    assert out[1].text == OCR_HEADER + "\nMerge sort splits the array in half"


def test_window_bars_repeated_across_screenshots_are_left_out():
    before = Block(seq=0, kind="paragraph", text="Screenshot", heading_path=["t", "Level 0"], loc=Location(),
                   extra={"ref": "#/texts/0"})
    content = {400: "ls -la shows hidden files", 401: "cat readme prints the file", 402: "file data shows its type"}
    pictures = [(f"#/pictures/{width}", Image.new("RGB", (width, 300)), ["#/texts/0"]) for width in content]

    def read(image):
        return [OcrLine("kali-linux VirtualBox [Running]", 0, 10.0, 0.9),     # the same window bar each time
                OcrLine(content[image.width], 40, 10.0, 0.9)]

    out, count = _add_picture_text([before], pictures, read, "t")
    assert count == 3
    assert [b.text.splitlines()[1:] for b in out[1:]] == [[text] for text in content.values()]


def test_lines_already_read_in_an_earlier_picture_and_hex_dumps_are_left_out():
    before = Block(seq=0, kind="paragraph", text="Screenshot", heading_path=["t", "Level 0"], loc=Location(),
                   extra={"ref": "#/texts/0"})
    banner = "This machine has a 64bit processor"                     # a login banner in every screenshot
    pictures = [("#/pictures/0", Image.new("RGB", (400, 300)), ["#/texts/0"]),
                ("#/pictures/1", Image.new("RGB", (401, 300)), ["#/texts/0"])]

    def read(image):
        last = "cat readme shows the file" if image.width == 400 else "grep millionth data.txt finds it"
        return [OcrLine(banner, 0, 10.0, 0.9), OcrLine("0020 - 08 f7 ea c1 f5 61 31 df", 0, 10.0, 0.9),
                OcrLine(last, 0, 10.0, 0.9)]

    out, count = _add_picture_text([before], pictures, read, "t")
    assert count == 2
    assert [b.text.splitlines()[1:] for b in out[1:]] == [
        [banner, "cat readme shows the file"], ["grep millionth data.txt finds it"]]


def test_a_stray_line_far_to_the_left_does_not_indent_all_the_others():
    lines = [OcrLine(f"line {n} of the terminal output", 600, 10.0, 0.9) for n in range(10)]
    lines.insert(5, OcrLine("-Wl,-z,norelro disable relro", 300, 10.0, 0.9))
    assert not any(line.startswith(" ") for line in _picture_text(lines).splitlines())