"""Track 1: digital PDF, DOCX and PPTX through Docling. Owner: Areeba.

The DOCX and PPTX files are built fresh in each test with python-docx and
python-pptx (both come with Docling), so no big files are committed. Reading a
PDF needs Docling's layout models, so the PDF rules are tested on a Docling
document built by hand; real PDFs run only in the slow test at the end.
"""
import re
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docling")      # installed from requirements.txt; until then these tests are skipped

import docx  # noqa: E402
from docling_core.types.doc import (BoundingBox, ContentLayer, DocItemLabel, DoclingDocument,  # noqa: E402
                                    ProvenanceItem, Size, TableCell, TableData)
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.ingestion import messages  # noqa: E402
from app.ingestion.contract import PrecheckResult, ResourceMeta  # noqa: E402
from app.ingestion.dispatcher import ingest_file  # noqa: E402
from app.ingestion.loaders import docs  # noqa: E402

SAMPLES = Path(__file__).parent / "samples" / "docs"


# ---------------------------------------------------------------------------
# Test files
# ---------------------------------------------------------------------------

@pytest.fixture
def normalization_docx(tmp_path) -> Path:
    d = docx.Document()
    d.add_heading("Normalization", level=0)                       # Word's "Title" style
    d.add_heading("First Normal Form", level=1)
    d.add_paragraph("A table is in 1NF when every cell holds one value.")
    d.add_paragraph("No repeating groups", style="List Bullet")
    d.add_heading("Example", level=2)
    table = d.add_table(rows=3, cols=2)
    for r, row in enumerate([("Student", "Course"), ("Ali", "DB"), ("Sara", "OOP")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_heading("Third Normal Form", level=1)
    d.add_paragraph("Remove transitive depen-\ndencies between non-key columns.")
    d.add_paragraph("Find the key", style="List Number")
    d.add_paragraph("Split the table", style="List Number")
    path = tmp_path / "normalization.docx"
    d.save(path)
    return path


@pytest.fixture
def stacks_pptx(tmp_path) -> Path:
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])                    # title slide
    s.shapes.title.text = "Stacks"
    s.placeholders[1].text = "Data Structures, Lecture 6"
    s = p.slides.add_slide(p.slide_layouts[1])                    # title and bullets
    s.shapes.title.text = "Push and pop"
    s.placeholders[1].text_frame.text = "push adds on top"
    s.notes_slide.notes_text_frame.text = "Ask what happens when you pop an empty stack."
    s = p.slides.add_slide(p.slide_layouts[5])                    # title and a table
    s.shapes.title.text = "Operations and cost"
    table = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for r, row in enumerate([("Operation", "Cost"), ("push", "O(1)"), ("pop", "O(1)")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    s = p.slides.add_slide(p.slide_layouts[6])                    # blank slide, no title
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Questions?"
    path = tmp_path / "stacks.pptx"
    p.save(path)
    return path


def lecture_pdf() -> DoclingDocument:
    """What Docling returns for a 4-page digital lecture PDF, built by hand (no models needed)."""
    d = DoclingDocument(name="lecture5")

    def on(page: int) -> ProvenanceItem:
        return ProvenanceItem(page_no=page, bbox=BoundingBox(l=50, t=50, r=500, b=80), charspan=(0, 0))

    for page in range(1, 5):
        d.add_page(page_no=page, size=Size(width=595, height=842))
    d.add_title("Lecture 5: Joins", prov=on(1))
    d.add_heading("Inner join", level=1, prov=on(1))
    d.add_text(DocItemLabel.TEXT, "An inner join keeps rows that match in both tables.", prov=on(1))
    d.add_heading("Example", level=2, prov=on(2))
    cells = [TableCell(text=value, start_row_offset_idx=r, end_row_offset_idx=r + 1,
                       start_col_offset_idx=c, end_col_offset_idx=c + 1, column_header=r == 0)
             for r, row in enumerate([("id", "name"), ("1", "Ali")]) for c, value in enumerate(row)]
    d.add_table(TableData(num_rows=2, num_cols=2, table_cells=cells), prov=on(2))
    d.add_heading("Outer join", level=1, prov=on(3))
    d.add_text(DocItemLabel.TEXT, "An outer join also keeps rows without a match.", prov=on(3))
    for page in range(1, 4):
        d.add_text(DocItemLabel.TEXT, f"Database Systems | Page {page}", prov=on(page))   # a footer Docling missed
        d.add_text(DocItemLabel.PAGE_FOOTER, str(page), prov=on(page),
                   content_layer=ContentLayer.FURNITURE)                                  # one it caught
    d.add_picture(prov=on(4))                                                             # page 4: a picture only
    return d


def pdf_meta() -> ResourceMeta:
    return ResourceMeta(resource_id="r-pdf", course_id="c1", file_name="lecture5.pdf", file_sha256="0" * 64,
                        precheck=PrecheckResult(source_type="pdf", pages_total=4))


def text_of(doc, start: str):
    return next(b for b in doc.blocks if b.text.startswith(start))


# ---------------------------------------------------------------------------
# DOCX (real Docling, fast)
# ---------------------------------------------------------------------------

def test_docx_headings_become_the_heading_path(normalization_docx, data_dir):
    doc = ingest_file(normalization_docx, "r-docx", "c1")
    assert doc.quality.status == "ok" and doc.parser == "docling"
    assert doc.title == "Normalization"
    assert text_of(doc, "A table is in 1NF").heading_path == ["Normalization", "First Normal Form"]
    assert text_of(doc, "Remove transitive").heading_path == ["Normalization", "Third Normal Form"]
    assert all(b.loc.page is None for b in doc.blocks)              # DOCX has no pages


def test_text_is_cleaned_but_tables_are_not(normalization_docx, data_dir):
    doc = ingest_file(normalization_docx, "r-docx", "c1")
    assert text_of(doc, "Remove transitive").text == "Remove transitive dependencies between non-key columns."
    tables = [b for b in doc.blocks if b.kind == "table"]
    assert len(tables) == 1
    assert tables[0].text.startswith("| Student") and "| Sara" in tables[0].text
    assert tables[0].heading_path == ["Normalization", "First Normal Form", "Example"]


def test_numbered_list_markers_are_kept(normalization_docx, data_dir):
    doc = ingest_file(normalization_docx, "r-docx", "c1")
    assert [b.extra.get("marker") for b in doc.blocks if b.text in ("Find the key", "Split the table")] == ["1.", "2."]


def test_table_cells_are_not_repeated_as_paragraphs(tmp_path, data_dir):
    d = docx.Document()
    d.add_heading("Marks", level=1)
    table = d.add_table(rows=2, cols=2)
    for c, value in enumerate(["Student", "Marks"]):
        table.cell(0, c).paragraphs[0].add_run(value).bold = True     # bold headers, as in most real tables
    table.cell(1, 0).text, table.cell(1, 1).text = "Ali", "18"
    path = tmp_path / "marks.docx"
    d.save(path)
    doc = ingest_file(path, "r-marks", "c1")
    assert [b.kind for b in doc.blocks if "Student" in b.text] == ["table"]


def test_word_footnotes_are_kept_but_page_headers_are_not(tmp_path, data_dir):
    """Docling files Word footnotes beside page headers and footers; only the footnotes are content."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import parse_xml

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Lecture notes, Fall 2026"
    d.add_heading("Normal forms", level=1)
    p = d.add_paragraph("Codd defined the first three normal forms.")
    footnotes = (f'<w:footnotes {w}>'
                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:id="1"><w:p><w:r><w:t>He published them in 1970 and 1971.</w:t></w:r></w:p></w:footnote>'
                 '</w:footnotes>')
    part = Part(PackURI("/word/footnotes.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                footnotes.encode("utf-8"), d.part.package)
    d.part.relate_to(part, RELATIONSHIP_TYPE.FOOTNOTES)
    p._p.append(parse_xml(f'<w:r {w}><w:footnoteReference w:id="1"/></w:r>'))
    path = tmp_path / "footnote.docx"
    d.save(path)
    doc = ingest_file(path, "r-footnote", "c1")
    assert any("published them in 1970" in b.text for b in doc.blocks)
    assert not any("Lecture notes" in b.text for b in doc.blocks)


# ---------------------------------------------------------------------------
# PPTX (real Docling, fast)
# ---------------------------------------------------------------------------

def test_pptx_path_is_deck_title_then_slide_title(stacks_pptx, data_dir):
    doc = ingest_file(stacks_pptx, "r-pptx", "c1")
    assert doc.quality.status == "ok" and doc.quality.units_total == 4
    assert all(b.loc.slide is not None for b in doc.blocks)
    assert {b.loc.slide: b.heading_path for b in doc.blocks} == {
        1: ["stacks", "Stacks"], 2: ["stacks", "Push and pop"],
        3: ["stacks", "Operations and cost"], 4: ["stacks", "Slide 4"]}


def test_speaker_notes_become_slide_notes(stacks_pptx, data_dir):
    doc = ingest_file(stacks_pptx, "r-pptx", "c1")
    notes = [b for b in doc.blocks if b.kind == "slide_notes"]
    assert [(b.loc.slide, b.text) for b in notes] == [(2, "Ask what happens when you pop an empty stack.")]


# ---------------------------------------------------------------------------
# Never twice, never crash
# ---------------------------------------------------------------------------

def test_the_same_file_is_converted_only_once(normalization_docx, data_dir, monkeypatch):
    first = ingest_file(normalization_docx, "r-first", "c1")

    def no_second_conversion():
        raise AssertionError("Docling ran again on a file it already read")

    monkeypatch.setattr(docs, "_converter", no_second_conversion)
    second = ingest_file(normalization_docx, "r-second", "c1")      # same file, uploaded again
    assert second.quality.status == "ok"
    assert [b.model_dump() for b in second.blocks] == [b.model_dump() for b in first.blocks]
    assert (data_dir / "r-first" / "raw" / "docling.json").exists()
    assert (data_dir / "r-second" / "raw" / "docling.json").exists()


def test_a_file_docling_cannot_read_fails_politely(tmp_path, data_dir):
    fake = tmp_path / "notes.docx"                     # a real zip, so it passes the pre-check, but not a Word file
    with zipfile.ZipFile(fake, "w") as z:
        z.writestr("hello.txt", "not a Word document")
    doc = ingest_file(fake, "r-fake", "c1")
    assert doc.quality.status == "failed"
    assert doc.quality.warnings == [messages.UNREADABLE]


# ---------------------------------------------------------------------------
# PDF rules, on a Docling document built by hand
# ---------------------------------------------------------------------------

def pdf_blocks():
    d = lecture_pdf()
    return docs.to_blocks(docs.walk(d, "pdf", docs.document_title(d, "pdf", "lecture5.pdf")), "pdf")


def test_pdf_options_keep_heading_levels_and_use_fast_tables():
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import TableFormerMode

    options = docs._converter().format_to_options[InputFormat.PDF].pipeline_options
    assert options.do_ocr is False                                       # scans go to the OCR loader
    assert options.heading_hierarchy_options.enabled                     # otherwise every heading is level 1
    assert options.table_structure_options.mode == TableFormerMode.FAST  # 1.2x faster on the 150-page book


def test_pdf_blocks_carry_their_page_and_heading_path():
    blocks = pdf_blocks()
    assert [(b.loc.page, b.kind, b.heading_path) for b in blocks] == [
        (1, "heading", ["Lecture 5: Joins"]),
        (1, "heading", ["Lecture 5: Joins", "Inner join"]),
        (1, "paragraph", ["Lecture 5: Joins", "Inner join"]),
        (2, "heading", ["Lecture 5: Joins", "Inner join", "Example"]),
        (2, "table", ["Lecture 5: Joins", "Inner join", "Example"]),
        (3, "heading", ["Lecture 5: Joins", "Outer join"]),
        (3, "paragraph", ["Lecture 5: Joins", "Outer join"]),
    ]
    assert [b.seq for b in blocks] == list(range(len(blocks)))



def test_labels_code_and_sentences_are_not_headings():
    for text in ("Example 3: // copy the arrays", "Sample output1:", "Task # 2", "Homework:", "int main()",
                 "cout << *yptr; // 10", "It is easier, however, to combine both definitions on one line:",
                 "will generate the function"):
        assert not docs._heading_like(text, "pdf"), text
    for text in ("Lecture 06", "2.8 The C++ string Class", "Mixing cin >> and cin.get", "Example",
                 "8.2 Focus on Problem Solving and Program Design: A Case Study", "string Member Functions and Operators",
                 "6.15 The exit() Function", "Solution: Virtual Inheritance"):
        assert docs._heading_like(text, "pdf"), text


def test_lectures_are_top_level_but_a_contents_list_is_not():
    def on(page: int) -> ProvenanceItem:
        return ProvenanceItem(page_no=page, bbox=BoundingBox(l=50, t=50, r=500, b=80), charspan=(0, 0))

    d = DoclingDocument(name="notes")
    for page, text in ((2, "Chapter 1 Intro"), (2, "Chapter 2 Loops"), (3, "Chapter 3 Arrays")):   # a contents page
        d.add_heading(text, level=2, prov=on(page))
    for page, lecture in ((5, "Lecture 01: Recap"), (9, "Lecture 02: Loops")):
        d.add_heading(lecture, level=3, prov=on(page))
        d.add_heading("Arrays", level=1, prov=on(page))
        d.add_heading("Example 1:", level=2, prov=on(page))
        d.add_text(DocItemLabel.TEXT, "int a[3];", prov=on(page))
    assert sorted(item.text for item, _ in d.iterate_items() if item.self_ref in docs._parts(d)) == [
        "Lecture 01: Recap", "Lecture 02: Loops"]                          # not "Chapter 3", the end of the list
    raw = docs.walk(d, "pdf", "notes")
    assert raw[-1]["path"] == ["notes", "Lecture 02: Loops", "Arrays"]       # "Example 1:" is text, not a heading
    assert [r["path"] for r in raw if r["text"] == "Chapter 3 Arrays"] == [["notes", "Chapter 3 Arrays"]]


def test_pdf_headers_and_footers_are_dropped():
    assert not any("Database Systems" in b.text or b.text.strip() in {"1", "2", "3"} for b in pdf_blocks())


def test_a_long_book_makes_one_table_writer(monkeypatch):
    """Making Docling's table writer checks the whole document, so doing it per table is slow on long books."""
    from docling_core.transforms.serializer import markdown

    made, real = [], markdown.MarkdownDocSerializer
    monkeypatch.setattr(markdown, "MarkdownDocSerializer", lambda **kw: made.append(1) or real(**kw))
    d = lecture_pdf()
    for page in (3, 4):
        cells = [TableCell(text=value, start_row_offset_idx=r, end_row_offset_idx=r + 1,
                           start_col_offset_idx=0, end_col_offset_idx=1, column_header=r == 0)
                 for r, value in enumerate(["step", str(page)])]
        d.add_table(TableData(num_rows=2, num_cols=1, table_cells=cells),
                    prov=ProvenanceItem(page_no=page, bbox=BoundingBox(l=50, t=50, r=500, b=80), charspan=(0, 0)))
    blocks = docs.to_blocks(docs.walk(d, "pdf", "Lecture 5: Joins"), "pdf")
    assert [b.loc.page for b in blocks if b.kind == "table"] == [2, 3, 4]
    assert next(b.text for b in blocks if b.kind == "table").startswith("|   id | name   |")
    assert len(made) == 1


def test_pdf_quality_report(monkeypatch, data_dir):
    monkeypatch.setattr(docs, "convert", lambda path, meta: (lecture_pdf(), True))
    doc = docs.load(Path("lecture5.pdf"), pdf_meta())
    assert doc.title == "Lecture 5: Joins"
    assert doc.quality.status == "ok" and doc.quality.units_total == 4
    assert doc.quality.warnings == [messages.no_text_units("Pages", [4])]
    assert doc.quality.chars_total == sum(len(b.text) for b in doc.blocks)


def test_a_paragraph_running_onto_the_next_page_counts_for_both_pages(monkeypatch, data_dir):
    def on(page: int) -> ProvenanceItem:
        return ProvenanceItem(page_no=page, bbox=BoundingBox(l=50, t=50, r=500, b=80), charspan=(0, 0))

    d = lecture_pdf()
    item = d.add_text(DocItemLabel.TEXT, "This paragraph starts on page 3 and ends on page 4.", prov=on(3))
    item.prov.append(on(4))                  # Docling joins the two halves into one item
    monkeypatch.setattr(docs, "convert", lambda path, meta: (d, True))
    doc = docs.load(Path("lecture5.pdf"), pdf_meta())
    assert doc.quality.warnings == []        # page 4 isn't reported as having no text
    block = text_of(doc, "This paragraph starts")
    assert block.loc.page == 3 and block.extra["pages"] == [3, 4]


def test_a_partly_read_file_is_marked_degraded(monkeypatch, data_dir):
    monkeypatch.setattr(docs, "convert", lambda path, meta: (lecture_pdf(), False))
    doc = docs.load(Path("lecture5.pdf"), pdf_meta())
    assert doc.quality.status == "degraded"
    assert messages.PARTIAL in doc.quality.warnings


# ---------------------------------------------------------------------------
# Your own sample files in samples/docs/ (PDFs need Docling's models: slow)
# ---------------------------------------------------------------------------

def _samples():
    files = sorted(p for p in SAMPLES.glob("*") if p.suffix.lower() in {".pdf", ".docx", ".pptx"})
    return [pytest.param(p, id=p.name, marks=[pytest.mark.slow] if p.suffix.lower() == ".pdf" else [])
            for p in files]


@pytest.mark.parametrize("path", _samples())
def test_my_sample_files(path, data_dir):
    doc = ingest_file(path, "sample-" + re.sub(r"[^A-Za-z0-9_-]+", "-", path.stem)[:50], "c1")
    assert doc.quality.status != "failed", doc.quality.warnings
    assert all(b.heading_path[0] == doc.title for b in doc.blocks)
    if doc.source_type == "pdf":
        assert all(1 <= b.loc.page <= doc.quality.units_total for b in doc.blocks)