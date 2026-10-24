import logging
import os
import re
import shutil
from collections import Counter, defaultdict, deque
from collections.abc import Callable
from functools import lru_cache
from pathlib import Path
from statistics import median
from typing import NamedTuple

from docx import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import nsmap, qn
from docx.text.paragraph import Paragraph
from docling_core.types.doc import ContentLayer, DocItem, DoclingDocument, PictureItem
from PIL.Image import Image

from app.ingestion import messages
from app.ingestion.cleaning.common import clean_text, detect_lang, title_from_filename
from app.ingestion.contract import Block, Location, ParsedDocument, ResourceMeta
from app.ingestion.converter import PARSER, PARSER_VERSION, convert
from app.ingestion.docling_blocks import to_blocks
from app.ingestion.documents import parsed_document
from app.ingestion.registry import register

logger = logging.getLogger("tta.ingestion")

_W = {"w": nsmap["w"]}
_TOC_STYLE = re.compile(r"toc (heading|\d)")                   # Word's names: "TOC Heading", "toc 1" ... "toc 9"
_SKIPPED_PARTS = {"Table of Contents", "Cover Pages"}          # Word content controls that are not course content
_P, _RUN, _HIDDEN = qn("w:p"), qn("w:r"), f'{qn("w:rPr")}/{qn("w:vanish")}'
_RUN_FONT, _RUN_STYLE = f'{qn("w:rPr")}/{qn("w:rFonts")}', f'{qn("w:rPr")}/{qn("w:rStyle")}'

_MONO_FONTS = {"consolas", "courier", "courier new", "lucida console", "lucida sans typewriter",
               "menlo", "monaco", "inconsolata", "andale mono"}
_MONO_WORDS = ("mono", "code")              # Cascadia Mono, Roboto Mono, Source Code Pro, Fira Code ...
_CODE_SHARE = 0.9                           # a paragraph is code when this share of its letters is monospace
_NOT_CODE_STYLE = re.compile(r"heading|title|caption|toc|list")
_RUN_BOLD, _RUN_SIZE = f'{qn("w:rPr")}/{qn("w:b")}', f'{qn("w:rPr")}/{qn("w:sz")}'
_OUTLINE = f'{qn("w:pPr")}/{qn("w:outlineLvl")}'
_TITLE_WORDS = 12                          # a bold title line is short
_TITLE_LEVELS = 3                          # bold titles nest at most this deep, bigger font = higher level
_CAPTION_START = re.compile(r"(fig(ure)?|table|chart|graph|diagram)\b\.?\s*[\d:]", re.IGNORECASE)

# Pictures: the text inside them is read with Tesseract OCR and kept as a figure_text block.
OCR_SWITCH = "DOCX_PICTURE_OCR"            # set to 0 to turn picture OCR off
OCR_HEADER = "Text read from a picture (OCR, may contain mistakes):"
_WINDOWS_TESSERACT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
_OCR_MIN_SIDE = 40                          # pixels: a smaller picture is an icon or a bullet
_OCR_MIN_WORDS = 3                          # fewer real words (3+ letters or digits): a photo or a logo
_OCR_EDGE_LINES = (4, 2)                    # top and bottom lines where window bars and status bars sit
_OCR_REPEAT_WORDS = 3                       # a line this long already read from an earlier picture is a repeat
_OCR_OUTLIER_CHARS = 20                     # a line starting this far left of the others is a stray, not code
_HEX_DUMP = re.compile(r"^\s*\S{3,4}\s+-\s+(?:\S{2}\s+){3}")  # "0020 - 08 f7 ea ...": a hex dump, garbled by OCR
_OCR_MAX_INDENT = 40                        # characters: further right is a pop-up or a status bar, not indentation
_OCR_MIN_SCORE = 0.4                        # Tesseract's own confidence; lines below it are left out
_OCR_LAYERS = {ContentLayer.BODY, ContentLayer.NOTES}
_BULLET_MISREAD = re.compile(r"^[¢«©°•·▪■●◦*>]\s+")      # how OCR reads a bullet point: written as "- "


class _WordFacts(NamedTuple):
    """What Word knows about its paragraphs but Docling does not tell us."""
    captions: Counter[str]          # paragraphs in the Caption style
    skipped: Counter[str]           # table of contents and cover page lines
    visible: dict[str, str]         # text of a paragraph with hidden words -> the same text without them
    code: list[list[str]]           # runs of paragraphs set in a monospace font, line by line as Word shows them
    titles: list[tuple[str, int]]   # bold title lines (key, level), only for a file without any heading style


def _norm(text: str) -> str:
    """Clean Word's text the same way to_blocks() cleans block text, so the two compare equal."""
    return " ".join(clean_text(text).split())


def _doc_parts(element) -> set[str]:
    """The kinds of Word content control around a paragraph, e.g. "Table of Contents" or "Cover Pages"."""
    return {
        kind
        for sdt in element.iterancestors(qn("w:sdt"))
        for kind in sdt.xpath("./w:sdtPr/w:docPartObj/w:docPartGallery/@w:val", namespaces=_W)
    }


def _is_hidden(run) -> bool:
    vanish = run.find(_HIDDEN)
    return vanish is not None and vanish.get(qn("w:val"), "true") not in ("0", "false")


def _own_runs(paragraph) -> list:
    """The paragraph's runs, also inside links and tracked insertions, but not those of a text box inside it."""
    return [run for run in paragraph.iter(_RUN) if next(run.iterancestors(_P)) is paragraph]


def _run_text(paragraph, keep_hidden: bool) -> str:
    """The paragraph's text, run by run. Deleted text is never included."""
    return "".join(run.text for run in _own_runs(paragraph) if keep_hidden or not _is_hidden(run))


def _is_mono(font: str | None) -> bool:
    name = (font or "").casefold()
    return bool(name) and (name in _MONO_FONTS or any(word in name for word in _MONO_WORDS))


def _style_font(style) -> str | None:
    """A style's font, looking through the styles it is based on."""
    while style is not None:
        if style.font.name:
            return style.font.name
        style = style.base_style
    return None


def _run_font(run, paragraph_font: str | None, styles_by_id: dict) -> str | None:
    """The font a run is shown in: its own, else its character style's, else the paragraph style's.

    The document's default font is left out on purpose: it says nothing about the author marking code.
    """
    fonts = run.find(_RUN_FONT)
    if fonts is not None and (name := fonts.get(qn("w:ascii")) or fonts.get(qn("w:hAnsi"))):
        return name
    char_style = run.find(_RUN_STYLE)
    if char_style is not None and (name := _style_font(styles_by_id.get(char_style.get(qn("w:val"))))):
        return name
    return paragraph_font


def _key(text: str) -> str:
    """Words only: a title still matches when a loose ":" or "-" around it was cleaned away."""
    return " ".join(word for word in _norm(text).split() if any(ch.isalnum() for ch in word))


def _is_title_line(raw: str, letters: int, mono: int) -> bool:
    """Short, one line, not ending like a sentence, not a caption and not code."""
    text = " ".join(raw.split())
    return (bool(letters) and mono < _CODE_SHARE * letters and "\n" not in raw.strip()
            and 0 < len(text.split()) <= _TITLE_WORDS and not text.endswith((".", ",", ";"))
            and not _CAPTION_START.match(text))


def _style_value(style, read):
    """A style's own value, else the value of the style it is based on, and so on."""
    while style is not None:
        if (value := read(style)) is not None:
            return value
        style = style.base_style
    return None


def _bold_share_and_size(paragraph, style, styles_by_id: dict) -> tuple[float, float]:
    """The share of letters and digits that are bold, and the biggest font size (points), of the visible runs."""
    paragraph_bold = bool(_style_value(style, lambda s: s.font.bold))
    paragraph_size = _style_value(style, lambda s: s.font.size)
    bold_letters = all_letters = 0
    size = 0.0
    for run in _own_runs(paragraph):
        if _is_hidden(run):
            continue
        letters = sum(ch.isalnum() for ch in run.text)
        if not letters:
            continue
        flag = run.find(_RUN_BOLD)
        if flag is not None:
            bold = flag.get(qn("w:val"), "true") not in ("0", "false")
        else:
            char_style = run.find(_RUN_STYLE)
            char_bold = _style_value(styles_by_id.get(char_style.get(qn("w:val"))), lambda s: s.font.bold) \
                if char_style is not None else None
            bold = paragraph_bold if char_bold is None else char_bold
        all_letters += letters
        bold_letters += letters if bold else 0
        run_size = run.find(_RUN_SIZE)
        points = int(run_size.get(qn("w:val"))) / 2 if run_size is not None else (
            paragraph_size.pt if paragraph_size is not None else 0.0)
        size = max(size, points)
    return (bold_letters / all_letters if all_letters else 0.0), size


def _word_facts(path: Path) -> _WordFacts:
    empty = _WordFacts(captions=Counter(), skipped=Counter(), visible={}, code=[], titles=[])
    try:
        word = WordDocument(str(path))
    except Exception:
        logger.warning("Word styles not read, python-docx failed on %s", path.name, exc_info=True)
        return empty
    captions, skipped, visible = Counter(), Counter(), {}
    code: list[list[str]] = []
    lines: list[str] = []                   # the code run being collected
    heading_mono = heading_letters = 0      # to tell a file typed all in Courier from one with code in it
    has_headings = False
    title_lines: list[tuple[str, float]] = []    # (key, font size) of bold title lines
    styles_by_id = {style.style_id: style for style in word.styles}

    def end_code_run() -> None:
        while lines and not lines[-1].strip():
            lines.pop()                     # a code run never ends in blank lines
        if lines:
            code.append(list(lines))
        lines.clear()

    # body.iter() reaches every paragraph, also inside content controls and tables.
    # word.paragraphs returns only the top-level ones and misses the table of contents and cover page.
    for element in word.element.body.iter(_P, qn("w:tbl")):
        if element.tag != _P:
            end_code_run()                  # a table ends a code run
            continue
        text = _norm(_run_text(element, keep_hidden=True))
        style = Paragraph(element, word).style
        style_name = style.name.casefold() if style is not None else ""
        if text:
            if style_name == "caption":
                captions[text] += 1
            elif _TOC_STYLE.fullmatch(style_name) or _doc_parts(element) & _SKIPPED_PARTS:
                skipped[text] += 1
            if any(_is_hidden(run) for run in _own_runs(element)):
                visible[text] = _norm(_run_text(element, keep_hidden=False))

        # Code and bold titles: only paragraphs in the body, not in a table cell or text box.
        # Code is a paragraph, not a heading, caption or list item, (nearly) all in a monospace font.
        if any(True for _ in element.iterancestors(qn("w:tc"), qn("w:txbxContent"))):
            continue
        raw = _run_text(element, keep_hidden=False)
        paragraph_font = _style_font(style)
        mono = total = 0
        for run in _own_runs(element):
            if _is_hidden(run):
                continue
            letters = len("".join(run.text.split()))
            total += letters
            if _is_mono(_run_font(run, paragraph_font, styles_by_id)):
                mono += letters
        if style_name.startswith(("heading", "title")):
            heading_mono, heading_letters = heading_mono + mono, heading_letters + total
        is_list = element.find(f'{qn("w:pPr")}/{qn("w:numPr")}') is not None
        if style_name.startswith(("heading", "title")) or element.find(_OUTLINE) is not None:
            has_headings = True
        elif not is_list and not _NOT_CODE_STYLE.search(style_name) and _is_title_line(raw, total, mono):
            bold, size = _bold_share_and_size(element, style, styles_by_id)
            if bold >= _CODE_SHARE:
                title_lines.append((_key(raw), size))
        if not raw.strip():
            if lines:
                lines.append("")            # a blank line inside a code run stays
            continue
        if (total and mono / total >= _CODE_SHARE and not is_list
                and not _NOT_CODE_STYLE.search(style_name)):
            lines.extend(raw.split("\n"))   # Shift+Enter inside a code paragraph starts a new line
        else:
            end_code_run()
    end_code_run()
    normal_font = _style_font(word.styles.default(WD_STYLE_TYPE.PARAGRAPH))
    if _is_mono(normal_font) or (heading_letters and heading_mono / heading_letters >= 0.5):
        code = []                           # the whole file is set in a monospace font: its look, not code
    titles: list[tuple[str, int]] = []
    if not has_headings:                    # a file with real headings keeps its bold lines as paragraphs
        sizes = sorted({size for _, size in title_lines}, reverse=True)
        titles = [(key, min(sizes.index(size) + 1, _TITLE_LEVELS)) for key, size in title_lines]
    return _WordFacts(captions=captions, skipped=skipped, visible=visible, code=code, titles=titles)


def _drop_word_extras(blocks: list[Block], skipped: Counter[str]) -> list[Block]:
    """Leave out the table of contents and the cover page, and take their headings out of the heading paths.

    Each skipped line is used once. Both come before the real content, so a real heading with the same
    words is kept.
    """
    out: list[Block] = []
    dropped: list[list[str]] = []          # heading paths of dropped headings, each ending in its own text
    for block in blocks:
        key = _norm(block.text)
        if block.kind in ("heading", "paragraph") and skipped[key] > 0:
            skipped[key] -= 1
            if block.kind == "heading":
                dropped.append(block.heading_path)
            continue
        remove = {len(d) - 1 for d in dropped if block.heading_path[:len(d)] == d}
        if remove:
            path = [h for i, h in enumerate(block.heading_path) if i not in remove]
            block = block.model_copy(update={"heading_path": path})
        out.append(block)
    return out


def _drop_hidden_text(blocks: list[Block], visible: dict[str, str]) -> list[Block]:
    """Word's hidden text (Font > Hidden) is not shown to students, so it is not course content."""
    out: list[Block] = []
    for block in blocks:
        key = _norm(block.text)
        if block.kind in ("heading", "paragraph") and key in visible:
            if not visible[key]:
                continue                                    # the whole paragraph was hidden
            block = block.model_copy(update={"text": visible[key]})
        out.append(block)
    return out


def _words(text: str) -> list[str]:
    """The words of a code text, line by line as the cleaner sees them. Bare symbols ("}", "|") are skipped:
    the cleaner drops a line holding nothing else, so they cannot be matched reliably."""
    return [word for line in text.split("\n") for word in _norm(line).split()
            if any(ch.isalnum() for ch in word)]


def _code_span(blocks: list[Block], first: int, want: list[str]) -> int:
    """How many blocks from `first` hold exactly the words of a code run, in order; 0 if they don't."""
    got: list[str] = []
    for count, block in enumerate(blocks[first:], start=1):
        if block.kind not in ("paragraph", "code"):
            return 0
        got += _words(block.text)
        if got != want[:len(got)]:
            return 0
        if len(got) == len(want):
            return count
    return 0


def _mark_code(blocks: list[Block], code: list[list[str]]) -> list[Block]:
    """Make each run of monospace paragraphs one code block, with Word's exact spacing.

    Docling marks a monospace paragraph as code only when it looks like code ({ } = < >, a call, a def line),
    and not at all in a file without a Normal style, so commands like "ls -la" come back as paragraphs with
    their spacing cleaned away. A run may come back as a mix of paragraph and code blocks, and a line the
    cleaner empties ("}") has no block at all, so the text is taken from Word.
    """
    out = list(blocks)
    start = 0
    for lines in code:
        want = _words("\n".join(lines))
        if not want:
            continue
        for first in range(start, len(out)):
            span = _code_span(out, first, want)
            if span:
                text = "\n".join(line.rstrip() for line in lines).strip("\n")
                out[first:first + span] = [out[first].model_copy(update={"kind": "code", "text": text, "lang": "code"})]
                start = first + 1
                break
    return out


def _mark_captions(blocks: list[Block], captions: Counter[str]) -> list[Block]:
    out: list[Block] = []
    for block in blocks:
        key = _norm(block.text)
        if block.kind == "paragraph" and captions[key] > 0:
            captions[key] -= 1
            block = block.model_copy(update={"kind": "caption"})
        out.append(block)
    return out


def _promote_titles(blocks: list[Block], titles: list[tuple[str, int]], title: str) -> list[Block]:
    """In a file without any heading style, its bold title lines are its headings: make them headings and
    give every block the heading path above it, so each section gets its own chunks.

    Without this, the chunker packs paragraphs by size alone and a title can end up at the bottom of the
    previous section's chunk, away from the text it names. A title with the same words as the file's title
    is already heading_path[0], so it adds no level.
    """
    if not titles or any(block.kind == "heading" for block in blocks):
        return blocks                       # Docling found real headings (e.g. a style in another language)
    levels: dict[str, deque[int]] = defaultdict(deque)
    for key, level in titles:
        levels[key].append(level)
    stack: list[tuple[int, str]] = []
    out: list[Block] = []
    for block in blocks:
        key = _key(block.text)
        if block.kind == "paragraph" and levels.get(key):
            level = levels[key].popleft()
            if block.text.casefold().split() != title.casefold().split():
                stack = [h for h in stack if h[0] < level] + [(level, block.text)]
            out.append(block.model_copy(update={"kind": "heading", "heading_path": [title] + [h[1] for h in stack]}))
        elif stack:
            out.append(block.model_copy(update={"heading_path": block.heading_path[:1] + [h[1] for h in stack]}))
        else:
            out.append(block)
    return out


class OcrLine(NamedTuple):
    """One line Tesseract read from a picture."""
    text: str
    left: int               # pixels from the picture's left edge to the line's first word
    char_width: float       # average pixels per character in the line
    score: float            # 0 to 1


Reader = Callable[[Image], list[OcrLine]]


@lru_cache(maxsize=1)
def _tesseract() -> Reader | None:
    """Tesseract through pytesseract, or None when either is missing (pictures are then skipped)."""
    try:
        import pytesseract
    except ImportError:
        logger.warning("picture OCR skipped: pytesseract is not installed")
        return None
    command = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
    if command is None and os.path.exists(_WINDOWS_TESSERACT):
        command = _WINDOWS_TESSERACT
    if command is None:
        logger.warning("picture OCR skipped: the tesseract program was not found")
        return None
    pytesseract.pytesseract.tesseract_cmd = command

    def read(image: Image) -> list[OcrLine]:
        data = pytesseract.image_to_data(image, config="--psm 6", output_type=pytesseract.Output.DICT)
        lines: dict[tuple[int, int, int], list[tuple[str, int, int, float]]] = {}
        for i, word in enumerate(data["text"]):
            score = float(data["conf"][i])
            if word.strip() and score >= 0:
                key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
                lines.setdefault(key, []).append((word, data["left"][i], data["width"][i], score / 100))
        out = []
        for words in lines.values():                     # Tesseract gives lines in reading order
            letters = sum(len(word) for word, *_ in words)
            out.append(OcrLine(
                text=" ".join(word for word, *_ in words), left=words[0][1],
                char_width=sum(width for _, _, width, _ in words) / max(letters, 1),
                score=sum(score for *_, score in words) / len(words),
            ))
        return out
    return read


def _ocr_reader() -> Reader | None:
    if os.environ.get(OCR_SWITCH, "1").strip() == "0":
        return None
    return _tesseract()


def _picture_text(lines: list[OcrLine]) -> str:
    """The picture's text, line by line. Indentation is rebuilt from where each line starts, because
    OCR loses it and Python code needs it; it is snapped to the smallest step so levels stay even."""
    kept = [line for line in lines if line.score >= _OCR_MIN_SCORE and line.text.strip()]
    if not kept:
        return ""
    width = median(line.char_width for line in kept) or 1.0
    middle = median(line.left for line in kept)
    near = [line.left for line in kept if (middle - line.left) / width <= _OCR_OUTLIER_CHARS]
    strays = len(kept) - len(near)
    left = min(near) if 0 < strays <= max(1, len(kept) // 10) else min(line.left for line in kept)
    indents = [max(0, round((line.left - left) / width)) for line in kept]
    indents = [indent if indent <= _OCR_MAX_INDENT else 0 for indent in indents]
    steps = [indent for indent in indents if indent >= 2]
    step = min(steps) if steps else 0
    return "\n".join(
        (" " * (round(indent / step) * step) if step and indent >= 2 else "") + _BULLET_MISREAD.sub("- ", line.text)
        for indent, line in zip(indents, kept)
    )


def _pictures(doc: DoclingDocument) -> list[tuple[str, Image | None, list[str]]]:
    """Every picture in reading order: its Docling id, its image, and the ids of the items before it."""
    out: list[tuple[str, Image | None, list[str]]] = []
    before: list[str] = []
    for item, _ in doc.iterate_items(traverse_pictures=True, included_content_layers=_OCR_LAYERS):
        if isinstance(item, PictureItem):
            try:
                image = item.get_image(doc) or (item.image.pil_image if item.image is not None else None)
            except Exception:                            # a picture Word can show but Pillow can't open (WMF)
                logger.warning("picture %s could not be opened", item.self_ref, exc_info=True)
                image = None
            out.append((item.self_ref, image, list(before)))
        elif isinstance(item, DocItem):
            before.append(item.self_ref)
    return out


def _real_words(text: str) -> int:
    """Words with at least 3 letters or digits. OCR of a logo or a photo gives scraps like "YZ J"."""
    return sum(1 for word in text.split() if sum(ch.isalnum() for ch in word) >= 3)


def _at_edge(k: int, count: int) -> bool:
    top, bottom = _OCR_EDGE_LINES
    return k < top or k >= count - bottom


def _bar_key(text: str) -> str:
    """A window bar's words only (3+ characters): OCR reads its icons and buttons differently each time.
    Digits are kept, so "ssh bandit0@..." and "ssh bandit1@..." stay different lines."""
    words = (re.sub(r"\W", "", word).lower() for word in text.split())
    return " ".join(word for word in words if len(word) >= 3)


def _add_picture_text(blocks: list[Block], pictures: list[tuple[str, Image | None, list[str]]],
                      read: Reader | None, title: str) -> tuple[list[Block], int]:
    """Put the text read from each picture right after the block before it, in the same section.

    Returns the blocks and how many pictures gave text. Skipped: pictures that are too small or can't be
    opened, and pictures with fewer than _OCR_MIN_WORDS real words (photos, logos). Left out: lines at the
    top or bottom that repeat in more than half the pictures (a screenshot's window title and menu bar),
    lines already read from an earlier picture (a login banner in every screenshot: kept the first time),
    and hex dumps.
    """
    if read is None or not pictures:
        return blocks, 0
    position = {block.extra.get("ref"): i for i, block in enumerate(blocks)}
    found: list[tuple[str, list[OcrLine], int]] = []       # picture id, its lines, the block before it
    for ref, image, before in pictures:
        if image is None or min(image.size) < _OCR_MIN_SIDE:
            continue
        lines = [line for line in read(image.convert("RGB")) if line.score >= _OCR_MIN_SCORE and line.text.strip()]
        index = next((position[r] for r in reversed(before) if r in position), -1)
        found.append((ref, lines, index))
    edge_counts: Counter[str] = Counter()
    for _, lines, _ in found:
        edge_counts.update({_bar_key(line.text) for k, line in enumerate(lines) if _at_edge(k, len(lines))} - {""})
    repeated = {key for key, n in edge_counts.items() if len(found) >= 3 and n > len(found) / 2}
    after: dict[int, list[Block]] = defaultdict(list)       # block index -> picture text blocks after it
    seen: set[str] = set()                                  # lines read from earlier pictures (a login banner)
    for ref, lines, index in found:
        kept = []
        for k, line in enumerate(lines):
            key = _bar_key(line.text)
            if _at_edge(k, len(lines)) and key in repeated:
                continue                                    # a window bar
            if key in seen and _real_words(line.text) >= _OCR_REPEAT_WORDS:
                continue                                    # already read from an earlier picture: kept there once
            if _HEX_DUMP.match(line.text):
                continue
            kept.append(line)
        seen.update(_bar_key(line.text) for line in lines if _real_words(line.text) >= _OCR_REPEAT_WORDS)
        text = _picture_text(kept)                          # indentation is measured after the bars are gone
        if _real_words(text) < _OCR_MIN_WORDS:
            continue
        path = blocks[index].heading_path if index >= 0 else [title]
        after[index].append(Block(
            seq=0, kind="figure_text", text=f"{OCR_HEADER}\n{text}", heading_path=path, loc=Location(),
            lang=detect_lang(text), extra={"ref": ref, "source": "ocr"},
        ))
    out = list(after.get(-1, []))
    for i, block in enumerate(blocks):
        out.append(block)
        out.extend(after.get(i, []))
    return out, sum(len(found_here) for found_here in after.values())


@register("docx", heavy=True)
def load(path: Path, meta: ResourceMeta) -> ParsedDocument:
    title = title_from_filename(meta.file_name)
    conversion = convert(path, "docx", meta.file_sha256)
    facts = _word_facts(path)
    blocks = to_blocks(conversion.document, "docx", title)
    blocks = _drop_word_extras(blocks, facts.skipped)
    blocks = _drop_hidden_text(blocks, facts.visible)
    blocks = _mark_code(blocks, facts.code)
    blocks = _mark_captions(blocks, facts.captions)
    blocks = _promote_titles(blocks, facts.titles, title)
    read = _ocr_reader()
    blocks, pictures_read = _add_picture_text(
        blocks, _pictures(conversion.document) if read else [], read, title)
    return parsed_document(
        meta, title=title, blocks=blocks, parser=PARSER, parser_version=PARSER_VERSION,
        units_total=len(blocks), problems=conversion.problems,
        notes=[messages.PICTURE_OCR] if pictures_read else [],
    )
